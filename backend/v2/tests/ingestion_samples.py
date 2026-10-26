"""Generated course-like fixtures, no real uploaded documents."""
from io import BytesIO


def text_pdf(pages=("Course Alpha", "Course Beta")):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    for value in pages:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"):NameObject("/Font"), NameObject("/Subtype"):NameObject("/Type1"),NameObject("/BaseFont"):NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"):writer._add_object(font)})})
        content=DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 50 700 Td ("+value.encode("ascii")+b") Tj ET")
        page[NameObject("/Contents")]=writer._add_object(content)
    result=BytesIO();writer.write(result)
    return result.getvalue()


def course_docx():
    from docx import Document
    document=Document()
    document.add_heading("课程标题",level=1)
    document.add_paragraph("段落内容与复杂度 O(n)")
    document.add_paragraph("列表条目",style="List Bullet")
    table=document.add_table(rows=2, cols=2)
    for row,values in zip(table.rows,[("算法","复杂度"),("查找","O(log n)")]):
        for cell,text in zip(row.cells,values):cell.text=text
    output=BytesIO();document.save(output)
    return output.getvalue()
